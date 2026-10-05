from pathlib import Path
import sys

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from docx_metadata import assert_docx_author, read_docx_author, set_docx_author  # noqa: E402


def test_docx_author_metadata_round_trip(tmp_path):
    output = tmp_path / "metadata.docx"
    document = set_docx_author(Document())
    document.add_paragraph("test")
    document.save(output)

    assert read_docx_author(output) == ("chenyechao", "chenyechao")
    assert_docx_author(output)


def test_docx_author_metadata_rejects_wrong_author(tmp_path):
    output = tmp_path / "wrong.docx"
    document = Document()
    document.core_properties.author = "someone-else"
    document.core_properties.last_modified_by = "someone-else"
    document.save(output)

    try:
        assert_docx_author(output)
    except ValueError as exc:
        assert "chenyechao" in str(exc)
    else:
        raise AssertionError("wrong author metadata was accepted")
