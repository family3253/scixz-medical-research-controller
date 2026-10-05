import json
import subprocess
import sys
from pathlib import Path

import yaml
from docx import Document
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"


def run_script(name, *args, check=True):
    return subprocess.run(
        [sys.executable, str(SCRIPTS / name), *map(str, args)],
        text=True,
        capture_output=True,
        check=check,
    )


def make_docx(path: Path, with_field=False):
    doc = Document()
    doc.core_properties.author = "Jane Doe"
    doc.core_properties.last_modified_by = "Jane Doe"
    doc.add_heading("A Test Manuscript", level=1)
    doc.add_paragraph("Methods and results remain unchanged.")
    doc.add_paragraph("Correspondence: jane.doe@example.edu")
    doc.add_paragraph("ORCID: 0000-0002-1825-0097")
    doc.add_paragraph("References")
    doc.add_paragraph("1. Doe J. Original title. Journal. 2020;1:1-2.")
    doc.add_paragraph("2. Roe R. Another title. Journal. 2021;2:3-4.")
    doc.sections[0].header.paragraphs[0].text = "Running header jane.doe@example.edu"
    doc.save(path)

    if with_field:
        import zipfile
        from xml.etree import ElementTree as ET

        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        with zipfile.ZipFile(path) as zf:
            files = {name: zf.read(name) for name in zf.namelist()}
        root = ET.fromstring(files["word/document.xml"])
        body = root.find("w:body", ns)
        p = ET.Element("{%s}p" % ns["w"])
        r = ET.SubElement(p, "{%s}r" % ns["w"])
        instr = ET.SubElement(r, "{%s}instrText" % ns["w"])
        instr.text = " ADDIN ZOTERO_ITEM CSL_CITATION {} "
        body.insert(1, p)
        files["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        tmp = path.with_suffix(".tmp.docx")
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, data in files.items():
                zf.writestr(name, data)
        tmp.replace(path)


def text_snapshot(path: Path):
    import zipfile
    from xml.etree import ElementTree as ET

    texts = []
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
        for el in root.iter():
            if el.tag.endswith("}t") and el.text:
                texts.append(el.text)
    return "\n".join(texts)


def write_requirements(path: Path, reference_style="vancouver"):
    data = {
        "schema_version": 1,
        "journal": "Example Journal",
        "article_type": "Original Article",
        "retrieved_at": "2026-09-21",
        "review_model": "single-anonymous",
        "source_records": [],
        "requirements": [
            {
                "id": "R001",
                "category": "layout",
                "rule": "A4 page, 1 inch margins, 1.5 spacing",
                "evidence": "test source",
                "applies_to": "Original Article",
                "source_url": "https://example.org/authors",
                "source_type": "official-journal",
                "accessed_at": "2026-09-21",
                "verification": "verified",
                "planned_change": "Apply layout actions",
                "change_class": "formatting",
                "status": "applied",
            }
        ],
        "actions": {
            "page_size": "A4",
            "margins_pt": {"top": 72, "right": 72, "bottom": 72, "left": 72},
            "default_font_name": "Arial",
            "default_font_size_pt": 11,
            "line_spacing": 1.5,
            "paragraph_spacing_after_pt": 6,
            "line_numbers": True,
            "page_numbers": True,
            "reference_style": reference_style,
            "preserve_dynamic_fields": True,
        },
        "conflicts": [],
        "unverified_items": [],
    }
    path.write_text(yaml.safe_dump(data, sort_keys=False), encoding="utf-8")


def test_template_starts_unverified():
    template = yaml.safe_load((ROOT / "assets" / "journal-requirements-template.yaml").read_text())
    item = template["requirements"][0]
    assert template["review_model"] == "unverified"
    assert item["verification"] == "unverified"
    assert item["source_type"] == "unverified"


def test_formatting_preserves_text_and_applies_explicit_actions(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "formatted.docx"
    requirements = tmp_path / "requirements.yaml"
    log = tmp_path / "changes.json"
    make_docx(source)
    write_requirements(requirements)
    before = text_snapshot(source)

    run_script("format_docx.py", source, output, "--requirements", requirements, "--change-log", log)

    after = text_snapshot(output)
    assert after == before
    doc = Document(output)
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.27
    assert round(section.page_height.inches, 2) == 11.69
    assert doc.styles["Normal"].font.name == "Arial"
    assert doc.styles["Normal"].font.size.pt == 11
    assert json.loads(log.read_text())["content_preserved"] is True


def test_reference_formatter_preserves_dynamic_fields(tmp_path):
    source = tmp_path / "fielded.docx"
    output = tmp_path / "fielded-out.docx"
    refs = tmp_path / "refs.json"
    make_docx(source, with_field=True)
    refs.write_text(json.dumps([
        {"id": 1, "formatted": "Doe J. Reformatted title. Journal. 2020;1:1-2."},
        {"id": 2, "formatted": "Roe R. Another title. Journal. 2021;2:3-4."},
    ]), encoding="utf-8")

    proc = run_script("reference_style.py", source, output, "--metadata", refs, "--style", "vancouver", check=False)
    assert proc.returncode != 0
    assert "dynamic" in (proc.stdout + proc.stderr).lower()
    assert not output.exists()


def test_reference_formatter_rewrites_plain_numeric_references(tmp_path):
    source = tmp_path / "plain.docx"
    output = tmp_path / "plain-out.docx"
    refs = tmp_path / "refs.json"
    make_docx(source)
    refs.write_text(json.dumps([
        {"id": 1, "formatted": "Doe J. Reformatted title. Journal. 2020;1:1-2."},
        {"id": 2, "formatted": "Roe R. Another title. Journal. 2021;2:3-4."},
    ]), encoding="utf-8")

    run_script("reference_style.py", source, output, "--metadata", refs, "--style", "vancouver")
    doc = Document(output)
    texts = [p.text for p in doc.paragraphs]
    assert "1. Doe J. Reformatted title. Journal. 2020;1:1-2." in texts
    assert "2. Roe R. Another title. Journal. 2021;2:3-4." in texts


def test_requirements_extractor_marks_heuristics_unverified(tmp_path):
    html = tmp_path / "instructions.html"
    output = tmp_path / "requirements.yaml"
    html.write_text(
        "<html><body><h1>Instructions for Authors</h1>"
        "<p>Abstract must be no more than 300 words.</p>"
        "<p>References should use Vancouver style.</p></body></html>",
        encoding="utf-8",
    )
    run_script("fetch_official_requirements.py", "--html-file", html, "--url", "https://example.org/authors", "--output", output, "--journal", "Example Journal")
    data = yaml.safe_load(output.read_text())
    assert data["requirements"]
    assert all(item["verification"] == "unverified" for item in data["requirements"])
    assert data["source_records"][0]["url"] == "https://example.org/authors"


def test_anonymize_removes_identity_markers_and_core_author(tmp_path):
    source = tmp_path / "source.docx"
    output = tmp_path / "blinded.docx"
    make_docx(source)
    run_script("anonymize_docx.py", source, output, "--term", "Jane Doe")
    scan = run_script("anonymity_scan.py", output, "--term", "Jane Doe")
    report = json.loads(scan.stdout)
    kinds = {item["kind"] for item in report["findings"]}
    assert "metadata_author" not in kinds
    assert "metadata_last_modified_by" not in kinds
    assert "email" not in kinds
    assert "orcid" not in kinds
    assert "user_term" not in kinds
