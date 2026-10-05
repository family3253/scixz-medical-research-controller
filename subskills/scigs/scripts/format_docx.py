from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

from common import content_digest, has_dynamic_fields, load_yaml, require_file


def all_paragraphs(doc):
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    return paragraphs


def set_style_font(style, name, size):
    if name:
        style.font.name = name
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for key in ("ascii", "hAnsi", "eastAsia"):
            rfonts.set(qn(f"w:{key}"), name)
    if size is not None:
        style.font.size = Pt(float(size))


def set_line_numbers(section):
    sect_pr = section._sectPr
    for existing in sect_pr.findall(qn("w:lnNumType")):
        sect_pr.remove(existing)
    element = OxmlElement("w:lnNumType")
    element.set(qn("w:countBy"), "1")
    element.set(qn("w:restart"), "newPage")
    sect_pr.append(element)


def set_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    if "PAGE" in paragraph._p.xml:
        return
    run = paragraph.add_run("Page ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def apply_actions(doc, actions: dict) -> list[str]:
    changes = []
    page_size = actions.get("page_size")
    if page_size:
        for section in doc.sections:
            if str(page_size).upper() == "A4":
                section.page_width, section.page_height = Mm(210), Mm(297)
            elif str(page_size).upper() in {"LETTER", "US LETTER"}:
                section.page_width, section.page_height = Inches(8.5), Inches(11)
            else:
                raise ValueError(f"Unsupported page_size action: {page_size}")
        changes.append(f"page_size={page_size}")

    margins = actions.get("margins_pt") or {}
    if any(value is not None for value in margins.values()):
        for section in doc.sections:
            for key, attr in (("top", "top_margin"), ("right", "right_margin"), ("bottom", "bottom_margin"), ("left", "left_margin")):
                if margins.get(key) is not None:
                    setattr(section, attr, Pt(float(margins[key])))
        changes.append(f"margins_pt={margins}")

    font_name = actions.get("default_font_name")
    font_size = actions.get("default_font_size_pt")
    if font_name or font_size is not None:
        for style in doc.styles:
            if style.type == 1:
                set_style_font(style, font_name, font_size)
        changes.append(f"default_font={font_name or 'unchanged'}:{font_size if font_size is not None else 'unchanged'}")

    line_spacing = actions.get("line_spacing")
    after = actions.get("paragraph_spacing_after_pt")
    if line_spacing is not None or after is not None:
        for paragraph in all_paragraphs(doc):
            if line_spacing is not None:
                paragraph.paragraph_format.line_spacing = float(line_spacing)
            if after is not None:
                paragraph.paragraph_format.space_after = Pt(float(after))
        changes.append(f"paragraph_spacing={line_spacing}:{after}")

    if actions.get("line_numbers") is True:
        for section in doc.sections:
            set_line_numbers(section)
        changes.append("line_numbers=enabled")
    if actions.get("page_numbers") is True:
        for section in doc.sections:
            set_page_number(section)
        changes.append("page_numbers=enabled")
    return changes


def main() -> int:
    parser = argparse.ArgumentParser(description="Apply explicit, content-preserving DOCX formatting actions")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--requirements", type=Path, required=True)
    parser.add_argument("--change-log", type=Path, required=True)
    args = parser.parse_args()

    source = require_file(args.input, "Input DOCX")
    requirements = load_yaml(args.requirements)
    actions = requirements.get("actions") or {}
    before = content_digest(source)
    if args.output.resolve() == source.resolve():
        raise SystemExit("Refusing to overwrite the source DOCX")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, args.output)
    doc = Document(args.output)
    doc.core_properties.author = "chenyechao"
    doc.core_properties.last_modified_by = "chenyechao"
    changes = apply_actions(doc, actions)
    doc.save(args.output)
    after = content_digest(args.output)
    result = {
        "source": str(source),
        "output": str(args.output),
        "content_digest_before": before,
        "content_digest_after": after,
        "content_preserved": before == after,
        "dynamic_fields_detected": has_dynamic_fields(source),
        "changes": changes,
        "status": "applied" if before == after else "blocked",
    }
    args.change_log.parent.mkdir(parents=True, exist_ok=True)
    args.change_log.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if before != after:
        args.output.unlink(missing_ok=True)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
