from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document

from common import has_dynamic_fields, require_file


NUMERIC_CITE_RE = re.compile(r"\[(\d+(?:\s*[-,;]\s*\d+)*)\]")
REFERENCE_RE = re.compile(r"^\s*(\d+)[.)]\s+")


def expand_citation(value: str) -> set[int]:
    values = set()
    for token in re.split(r"[,;]", value):
        token = token.strip()
        if "-" in token:
            left, right = [int(x.strip()) for x in token.split("-", 1)]
            values.update(range(min(left, right), max(left, right) + 1))
        elif token.isdigit():
            values.add(int(token))
    return values


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit numeric citation-to-reference integrity")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    path = require_file(args.docx, "DOCX")
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    cited = set()
    for match in NUMERIC_CITE_RE.finditer(text):
        cited.update(expand_citation(match.group(1)))
    refs = set()
    in_refs = False
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().casefold() in {"references", "bibliography", "reference list"}:
            in_refs = True
            continue
        if in_refs:
            match = REFERENCE_RE.match(paragraph.text)
            if match:
                refs.add(int(match.group(1)))
    missing = sorted(cited - refs)
    uncited = sorted(refs - cited)
    result = {
        "file": str(path),
        "dynamic_fields_detected": has_dynamic_fields(path),
        "cited_numeric_ids": sorted(cited),
        "bibliography_ids": sorted(refs),
        "missing_bibliography_entries": missing,
        "uncited_bibliography_entries": uncited,
        "status": "blocked" if missing else ("manual_action_required" if uncited else "pass"),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 2 if missing else 0


if __name__ == "__main__":
    raise SystemExit(main())
